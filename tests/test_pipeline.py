from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from app.catalogue.loader import CatalogueLoader
from app.catalogue.matcher import CatalogueMatcher
from app.clustering.engine import ClusteringEngine
from app.config.settings import DemoSettings
from app.pdf.analyzer import PDFAnalyzer
from app.rules.discovery import RuleDiscoveryEngine
from app.services.demo_runner import DemoRunner


def test_catalogue_loading():
    df, profile = CatalogueLoader().load(DemoSettings().catalogue_path)
    assert "PartNumber" in df.columns
    assert profile.row_count > 100
    assert "Technical attribute 1" in profile.technical_columns


def test_pdf_loading_and_part_numbers():
    pdf_path = next(path for path in DemoSettings().pdf_paths if path.exists())
    analysis = PDFAnalyzer().analyze(pdf_path)
    assert analysis.page_count >= 1
    assert analysis.part_numbers
    assert analysis.fields["Drawing number"] not in {"", "UNKNOWN", "NOT_FOUND"}


def test_catalogue_matching():
    settings = DemoSettings()
    df, _ = CatalogueLoader().load(settings.catalogue_path)
    analyses = [PDFAnalyzer().analyze(path) for path in settings.pdf_paths]
    matches = CatalogueMatcher().match(df, analyses)
    assert matches
    assert any(match.status in {"MATCH", "PREDICTED MATCH", "AMBIGUOUS"} for match in matches)


def test_rule_discovery_and_configuration_matrix():
    df, _ = CatalogueLoader().load(DemoSettings().catalogue_path)
    results = RuleDiscoveryEngine().discover(df)
    assert results["influence"]
    assert "configuration_matrix" in results


def test_clustering():
    df, _ = CatalogueLoader().load(DemoSettings().catalogue_path)
    clusters = ClusteringEngine().run(df)
    assert clusters["points"]
    assert clusters["explanations"]


def test_demo_exports(tmp_path):
    settings = DemoSettings(output_dir=tmp_path)
    progress = []
    result, generated, evidence = DemoRunner(settings, progress_callback=lambda percent, message: progress.append((percent, message))).run()
    assert len(generated) >= 2
    assert len(evidence) >= 2
    assert result.output_files["catalogue_csv"].exists()
    assert result.output_files["report"].exists()
    assert progress[0][0] == 0
    assert progress[-1][0] == 100
    assert [p for p, _ in progress] == sorted(p for p, _ in progress)
    workbook = pd.ExcelFile(result.output_files["catalogue_xlsx"])
    assert workbook.sheet_names[0] == "BRAKES_LEAN-CATALOGUE V2_TECH"
    for expected_sheet in {
        "SUMMARY",
        "PRODUCT CATALOGUE",
        "DIMENSIONS",
        "TECHNICAL PARAMETERS",
        "BOM",
        "TORQUE & FASTENERS",
        "REVISION HISTORY",
        "EXTRACTION SOURCES",
        "QUALITY REVIEW",
        "RAW DATA",
    }:
        assert expected_sheet in workbook.sheet_names
    styled_workbook = load_workbook(result.output_files["catalogue_xlsx"], read_only=False)
    assert styled_workbook["DIMENSIONS"].tables
    assert styled_workbook["DIMENSIONS"].freeze_panes == "A5"


def test_generated_catalogue_keeps_template_structure_and_pdf_codes(tmp_path):
    settings = DemoSettings(output_dir=tmp_path)
    template, _ = CatalogueLoader().load(settings.catalogue_path)
    _, generated, _ = DemoRunner(settings).run()
    assert list(generated.columns) == list(template.columns)

    generated_parts = set(generated["PartNumber"])
    for expected_part in {
        "FT0024835-100",
        "FT0129531-100",
        "FT0105784-100",
        "FT0027389-100",
        "FT0127469-100",
        "FT0127470-100",
    }:
        assert expected_part in generated_parts


def test_generated_catalogue_maps_additional_pdf_characteristics(tmp_path):
    settings = DemoSettings(output_dir=tmp_path)
    _, generated, _ = DemoRunner(settings).run()
    by_part = generated.set_index("PartNumber")

    assert by_part.loc["FT0027389-100", "Technical attribute 1"] == "TYPE 20 - 120°"
    assert by_part.loc["FT0127469-100", "Technical attribute 1"] == 'OUTLET 3/4" GAS UNI-ISO 228'
    assert by_part.loc["FT0127469-100", "Technical attribute 2"] == "ENVELOPE 500 x 370 x 265 mm"
    assert "WORKING PRESSURE 11 bar(g)" in by_part.loc["FT0127470-100", "Technical attribute 3"]


def test_generated_workbook_contains_aggressive_technical_characteristics(tmp_path):
    settings = DemoSettings(output_dir=tmp_path)
    result, _, evidence = DemoRunner(settings).run()
    catalogue = pd.read_excel(result.output_files["catalogue_xlsx"], sheet_name="BRAKES_LEAN-CATALOGUE V2_TECH", header=3).drop(columns=["Unnamed: 0"], errors="ignore").fillna("")

    assert "Technical attribute 12" in catalogue.columns
    assert len(evidence[evidence["extraction_method"].eq("technical_miner_tfidf")]) >= 150
    assert any(catalogue["Technical attribute 5"].astype(str).str.contains("IP54|IP 54", regex=True))
    assert any(catalogue["Technical attribute 9"].astype(str).str.contains("STAINLESS STEEL|stainless steel", regex=True))
    assert any(catalogue["Technical attribute 12"].astype(str).str.contains("EN 837-1|ISO 8573-1", regex=True))


def test_extended_catalogue_promotes_mined_manometer_characteristics(tmp_path):
    settings = DemoSettings(output_dir=tmp_path)
    result, generated, _ = DemoRunner(settings).run()
    workbook_catalogue = pd.read_excel(result.output_files["catalogue_xlsx"], sheet_name="BRAKES_LEAN-CATALOGUE V2_TECH", header=3).drop(columns=["Unnamed: 0"], errors="ignore").fillna("")
    by_part = workbook_catalogue.set_index("PartNumber")
    main_by_part = generated.set_index("PartNumber")

    assert "Technical attribute 12" in workbook_catalogue.columns
    assert "CLASS 1.0" in by_part.loc["FT0105784-100", "Technical attribute 5"]
    assert "IP 54" in by_part.loc["FT0105784-100", "Technical attribute 5"]
    assert "3-4-3 ISO 8573-1" in by_part.loc["FT0105784-100", "Technical attribute 8"]
    assert "BLOW-OUT 13" in by_part.loc["FT0105784-100", "Technical attribute 11"]
    assert "EN 837-1" in by_part.loc["FT0105784-100", "Technical attribute 12"]
    assert main_by_part.loc["FT0105784-100", "Technical attribute 1"] == "DIAMETER 80"


def test_engineering_model_extracts_structured_entities_from_drawings():
    settings = DemoSettings()
    analyses = {path.name: PDFAnalyzer().analyze(path) for path in settings.pdf_paths if path.exists()}

    gauge = analyses["1-498149_B03.pdf"].engineering
    assert gauge is not None
    assert any(d.dimension_type == "thread" and d.thread_designation == "M16X1.5" for d in gauge.dimensions)
    assert any(p.name == "Operating temperature" and p.value_min == -50 and p.value_max == 60 for p in gauge.parameters)
    assert any(s.standard == "DIN EN 837-1" for s in gauge.standards)
    assert any(m.material == "STAINLESS STEEL" for m in gauge.materials)

    compressor = analyses["FT0127469-100.pdf"].engineering
    assert compressor is not None
    assert any(d.nominal_value == 135 and d.upper_tolerance == 0.5 for d in compressor.dimensions)
    assert any(t.torque and t.unit == "Nm" for t in compressor.torque_requirements)
    assert compressor.drawing_references
    assert compressor.revisions
    assert compressor.raw_data
    assert any(source.source.bbox for source in compressor.dimensions if source.source.method.startswith("NATIVE"))
    assert compressor.inspections[0].has_text_layer
