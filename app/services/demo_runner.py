from __future__ import annotations

from copy import copy
import json
from pathlib import Path
import re
from typing import Any, Callable

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

from app.anomalies.detector import AnomalyDetector
from app.bom.analyzer import BOMAnalyzer
from app.catalogue.loader import CatalogueLoader
from app.catalogue.matcher import CatalogueMatcher
from app.clustering.engine import ClusteringEngine
from app.config.settings import DemoSettings, ensure_directories
from app.models.domain import DemoRunResult, Evidence
from app.pdf.analyzer import PDFAnalyzer
from app.reporting.report import ReportBuilder
from app.rules.discovery import RuleDiscoveryEngine
from app.similarity.engine import SimilarityEngine
from app.visualization.charts import ChartBuilder


class DemoRunner:
    def __init__(self, settings: DemoSettings | None = None, progress_callback: Callable[[int, str], None] | None = None) -> None:
        self.settings = settings or DemoSettings()
        self.progress_callback = progress_callback

    def run(self) -> tuple[DemoRunResult, pd.DataFrame, pd.DataFrame]:
        pdf_paths = [path for path in self.settings.pdf_paths if path.exists()]
        total_steps = 13 + len(pdf_paths)
        completed_steps = 0

        def advance(message: str) -> None:
            nonlocal completed_steps
            completed_steps += 1
            percent = min(100, round((completed_steps / total_steps) * 100))
            self._progress(percent, message)

        self._progress(0, "Preparing analysis")
        ensure_directories()
        output_dir = self.settings.output_dir
        output_dir.mkdir(parents=True, exist_ok=True)
        advance("Prepared output folders")

        catalogue, profile = CatalogueLoader().load(self.settings.catalogue_path)
        advance("Loaded sample catalogue as schema/reference template")
        pdf_analyzer = PDFAnalyzer()
        analyses = []
        for path in pdf_paths:
            analyses.append(pdf_analyzer.analyze(path))
            advance(f"Analyzed PDF: {path.name}")

        evidence = self._build_evidence_frame(analyses)
        advance("Built extraction evidence")
        matches = CatalogueMatcher().match(catalogue, analyses)
        advance("Matched PDFs to catalogue")
        generated = self._build_generated_catalogue(catalogue, analyses, matches)
        analysis_catalogue = generated.copy()
        demo_focus = self._build_demo_focus(catalogue, analysis_catalogue, expected_count=len(pdf_paths))
        advance("Generated catalogue rows from PDF drawings")
        rule_results = RuleDiscoveryEngine().discover(analysis_catalogue)
        advance("Discovered PDF catalogue PartNumber rules")
        clusters = ClusteringEngine().run(analysis_catalogue, method="kmeans", n_clusters=5)
        advance("Computed PDF catalogue clusters and PCA projection")
        similarity = SimilarityEngine().find_similar(analysis_catalogue)
        advance("Computed technical similarity")
        anomalies = AnomalyDetector().detect(analysis_catalogue, analyses, rule_results)
        advance("Detected anomalies")
        bom_summary = BOMAnalyzer().summarize(analyses)
        advance("Summarized BOM/component reuse")

        output_files: dict[str, Path] = {}
        output_files["catalogue_csv"] = output_dir / "catalogue_generated.csv"
        output_files["catalogue_xlsx"] = output_dir / "catalogue_generated.xlsx"
        output_files["evidence_csv"] = output_dir / "extraction_evidence.csv"
        output_files["analysis_json"] = output_dir / "analysis_results.json"
        output_files["anomalies_csv"] = output_dir / "anomalies.csv"
        output_files["bom_csv"] = output_dir / "bom_components.csv"
        output_files["focus_catalogue_csv"] = output_dir / "focus_catalogue.csv"
        output_files["focus_catalogue_xlsx"] = output_dir / "focus_catalogue.xlsx"
        client_catalogue = self._build_client_catalogue(generated, analyses)
        client_catalogue.to_csv(output_files["catalogue_csv"], index=False)
        self._write_catalogue_workbook(output_files["catalogue_xlsx"], generated, analyses, matches, evidence, anomalies)
        analysis_catalogue.to_csv(output_files["focus_catalogue_csv"], index=False)
        analysis_catalogue.to_excel(output_files["focus_catalogue_xlsx"], index=False)
        evidence.to_csv(output_files["evidence_csv"], index=False)
        pd.DataFrame(anomalies).to_csv(output_files["anomalies_csv"], index=False)
        pd.DataFrame([row for a in analyses for row in a.bom_rows]).to_csv(output_files["bom_csv"], index=False)
        self._write_json(output_files["analysis_json"], analyses, matches, rule_results, anomalies, clusters, similarity, bom_summary, demo_focus)
        advance("Exported CSV, XLSX, and JSON outputs")
        bom_rows = [row for a in analyses for row in a.bom_rows]
        chart_builder = ChartBuilder()
        visual_title = demo_focus["filter_value"]
        chart_paths = chart_builder.build_all(analysis_catalogue, clusters, output_dir, anomalies=anomalies, evidence=evidence, bom_rows=bom_rows, rule_results=rule_results, analyses=analyses)
        chart_paths.update(chart_builder.build_demo_focus(analysis_catalogue, clusters, output_dir, len(analysis_catalogue), visual_title))
        output_files.update(chart_paths)
        output_files.update(chart_builder.build_visual_packs(output_files, output_dir, visual_title))
        advance("Generated PDF catalogue visual analytics")

        result = DemoRunResult(
            catalogue_profile=profile,
            pdf_analyses=analyses,
            generated_catalogue_rows=len(generated),
            evidence_rows=len(evidence),
            demo_focus=demo_focus,
            matches=matches,
            rule_results=rule_results,
            anomalies=anomalies,
            clusters=clusters,
            similarity=similarity,
            output_files=output_files,
        )
        output_files["report"] = ReportBuilder().build(result, analysis_catalogue, output_dir / "DEMO_ANALYSIS_REPORT.md", bom_summary)
        advance("Built analysis report")
        self._progress(100, "Analysis complete")
        return result, generated, evidence

    def _progress(self, percent: int, message: str) -> None:
        if self.progress_callback:
            self.progress_callback(percent, message)

    def _build_generated_catalogue(self, catalogue: pd.DataFrame, analyses: list[Any], matches: list[Any]) -> pd.DataFrame:
        rows: list[dict[str, Any]] = []
        catalogue_columns = list(catalogue.columns)
        match_by_pdf = {match.source_pdf: match for match in matches}
        for analysis in analyses:
            match = match_by_pdf.get(analysis.source_pdf.name)
            row = self._matched_catalogue_base(catalogue, match) if match else {column: "" for column in catalogue.columns}
            row["OPS Product Category"] = row.get("OPS Product Category") or "C - BRAKE CONTROL"
            has_catalogue_match = bool(match and match.matched_part_number)
            is_variant_reference = self._is_variant_reference(analysis, match)
            if has_catalogue_match and not is_variant_reference:
                row["Product Family"] = self._existing_or_known(row.get("Product Family"), analysis.fields.get("Product Family"))
                row["Product Name"] = self._existing_or_known(row.get("Product Name"), analysis.fields.get("Product Name"))
                row["Product Type"] = self._existing_or_known(row.get("Product Type"), analysis.fields.get("Product Type"))
            else:
                row["Product Family"] = self._known_or_existing(analysis.fields.get("Product Family"), row.get("Product Family"))
                row["Product Name"] = self._known_or_existing(analysis.fields.get("Product Name"), row.get("Product Name"))
                row["Product Type"] = self._known_or_existing(analysis.fields.get("Product Type"), row.get("Product Type"))
            self._map_technical_attributes(row, analysis, prefer_existing=has_catalogue_match and not is_variant_reference)
            configuration = analysis.fields.get("Configuration", "")
            if not row.get("Configurability sheet") or is_variant_reference:
                row["Configurability sheet"] = configuration if configuration and configuration != "UNKNOWN" else "PDF - REVIEW"
            part_number, master_pn = self._catalogue_part_numbers(analysis, match, is_variant_reference)
            row["PartNumber"] = part_number
            row["Master PN"] = master_pn
            row["Maturity"] = row.get("Maturity") or analysis.fields.get("Status") or "Extracted"
            row["Preferred"] = row.get("Preferred") or "Needs review"
            row["Q.,ty in 2026"] = row.get("Q.,ty in 2026", "")
            rows.append({column: row.get(column, "") for column in catalogue_columns})
        return pd.DataFrame(rows, columns=catalogue_columns)

    def _is_variant_reference(self, analysis: Any, match: Any) -> bool:
        drawing = str(analysis.fields.get("Drawing number", ""))
        return bool(match and match.matched_part_number and drawing.startswith("FT") and match.matched_part_number != drawing)

    def _catalogue_part_numbers(self, analysis: Any, match: Any, is_variant_reference: bool) -> tuple[str, str]:
        drawing = str(analysis.fields.get("Drawing number", "UNKNOWN"))
        if is_variant_reference and drawing.startswith("FT"):
            return drawing, drawing
        if match and match.matched_part_number:
            return match.matched_part_number, match.matched_master_pn or match.matched_part_number
        part_number = self._primary_part(analysis.part_numbers, drawing)
        return part_number, self._master_part(analysis.part_numbers, drawing)

    def _analysis_columns(self, analysis: Any, match: Any) -> dict[str, Any]:
        evidence = analysis.evidence
        avg_confidence = round(sum(ev.confidence for ev in evidence) / max(1, len(evidence)), 2)
        review_fields = sorted(field for field, value in analysis.fields.items() if value in {"UNKNOWN", "NOT_FOUND", "NEEDS REVIEW"})
        return {
            "Source PDF": analysis.source_pdf.name,
            "PDF Page Count": analysis.page_count,
            "Catalogue Match Status": match.status if match else "UNMATCHED",
            "Catalogue Match Reason": match.reason if match else "No match object",
            "Catalogue Candidate Count": match.candidate_count if match else 0,
            "Matched PartNumber": match.matched_part_number if match else "",
            "Matched Master PN": match.matched_master_pn if match else "",
            "Drawing number": analysis.fields.get("Drawing number", ""),
            "Drawing revision": analysis.fields.get("Drawing revision", ""),
            "Drawing title": analysis.fields.get("Drawing title", ""),
            "Extracted diameter": analysis.fields.get("Diameter", ""),
            "Extracted LED": analysis.fields.get("LED", ""),
            "Extracted pressure": analysis.fields.get("Pressure", ""),
            "Extracted mounting": analysis.fields.get("Mounting", ""),
            "Extracted configuration": analysis.fields.get("Configuration", ""),
            "Extracted status": analysis.fields.get("Status", ""),
            "Extracted drain": analysis.fields.get("Drain", ""),
            "Extracted contact": analysis.fields.get("Contact", ""),
            "Extracted handle": analysis.fields.get("Handle", ""),
            "Extracted fitting": analysis.fields.get("Fitting", ""),
            "Extracted accuracy class": analysis.fields.get("Accuracy class", ""),
            "Extracted protection degree": analysis.fields.get("Protection degree", ""),
            "Extracted electrical rating": analysis.fields.get("Electrical rating", ""),
            "Extracted illumination detail": analysis.fields.get("Illumination detail", ""),
            "Extracted air quality": analysis.fields.get("Compressed air quality", ""),
            "Extracted case material": analysis.fields.get("Case material", ""),
            "Extracted frame ring": analysis.fields.get("Frame ring", ""),
            "Extracted pointer": analysis.fields.get("Pointer", ""),
            "Extracted pointer system": analysis.fields.get("Pointer system", ""),
            "Extracted process connection": analysis.fields.get("Process connection", ""),
            "Extracted lens": analysis.fields.get("Lens", ""),
            "Extracted restrictor screw": analysis.fields.get("Restrictor screw", ""),
            "Extracted safety blow-out": analysis.fields.get("Safety blow-out", ""),
            "Extracted standards": analysis.fields.get("Standards", ""),
            "Extracted material finishes": analysis.fields.get("Material finishes", ""),
            "Extracted outlet connection": analysis.fields.get("Outlet connection", ""),
            "Extracted envelope dimensions": analysis.fields.get("Envelope dimensions", ""),
            "Extracted mounting holes": analysis.fields.get("Mounting holes", ""),
            "Extracted key slot": analysis.fields.get("Key slot", ""),
            "Extracted working pressure": analysis.fields.get("Working pressure", ""),
            "Extracted working temperature": analysis.fields.get("Working temperature", ""),
            "Extracted startup temperature": analysis.fields.get("Startup temperature", ""),
            "Extracted duty cycle": analysis.fields.get("Duty cycle", ""),
            "Extracted starts per hour": analysis.fields.get("Starts per hour", ""),
            "Extracted weight": analysis.fields.get("Weight", ""),
            "Detected identifiers": ", ".join(analysis.part_numbers),
            "Detected variants": ", ".join(analysis.variants),
            "BOM row count": len(analysis.bom_rows),
            "BOM component identifiers": ", ".join(str(row.get("ComponentPartNumber", "")) for row in analysis.bom_rows[:30]),
            "Average extraction confidence": avg_confidence,
            "Fields needing review": ", ".join(review_fields),
            "PDF warnings": "; ".join(analysis.warnings),
            "Analysis note": "Generated from PDF extraction; sample Excel used only as schema/reference.",
        }

    def _write_catalogue_workbook(self, path: Path, generated: pd.DataFrame, analyses: list[Any], matches: list[Any], evidence: pd.DataFrame, anomalies: list[dict[str, Any]]) -> None:
        client_catalogue = self._build_client_catalogue(generated, analyses)
        self._write_template_catalogue_sheet(path, client_catalogue)
        wb = load_workbook(path)
        frames = self._engineering_workbook_frames(client_catalogue, analyses, matches, evidence, anomalies)
        for name in frames:
            if name in wb.sheetnames:
                del wb[name]
        for sheet_name, frame in frames.items():
            ws = wb.create_sheet(sheet_name)
            if sheet_name == "SUMMARY":
                self._write_summary_sheet(ws, frame)
            else:
                self._write_dataframe_sheet(ws, sheet_name, frame)
        wb.save(path)

    def _engineering_workbook_frames(self, client_catalogue: pd.DataFrame, analyses: list[Any], matches: list[Any], evidence: pd.DataFrame, anomalies: list[dict[str, Any]]) -> dict[str, pd.DataFrame]:
        return {
            "SUMMARY": self._summary_frame(analyses, evidence, anomalies),
            "PRODUCT CATALOGUE": client_catalogue,
            "PRODUCTS": self._products_frame(analyses, matches),
            "VARIANTS": self._variants_frame(analyses),
            "BOM": self._bom_frame(analyses),
            "COMPONENTS": self._components_frame(analyses),
            "DIMENSIONS": self._dimensions_frame(analyses),
            "TECHNICAL PARAMETERS": self._parameters_frame(analyses),
            "MATERIALS": self._materials_frame(analyses),
            "STANDARDS": self._standards_frame(analyses),
            "TORQUE & FASTENERS": self._torque_fasteners_frame(analyses),
            "ASSEMBLIES": self._assemblies_frame(analyses),
            "DRAWING VIEWS": self._views_frame(analyses),
            "DRAWING REFERENCES": self._references_frame(analyses),
            "CONNECTIONS": self._connections_frame(analyses),
            "SCHEMATICS": self._schematics_frame(analyses),
            "NOTES & INSTRUCTIONS": self._notes_frame(analyses),
            "REVISION HISTORY": self._revisions_frame(analyses),
            "IDENTIFICATION": self._identification_frame(analyses),
            "EXTRACTION SOURCES": self._sources_frame(analyses),
            "QUALITY REVIEW": self._quality_frame(analyses, anomalies),
            "RAW DATA": self._raw_data_frame(analyses),
        }

    def _summary_frame(self, analyses: list[Any], evidence: pd.DataFrame, anomalies: list[dict[str, Any]]) -> pd.DataFrame:
        rows: list[dict[str, Any]] = []
        totals: dict[str, int] = {
            "PDF drawings analyzed": len(analyses),
            "Pages processed": sum(analysis.page_count for analysis in analyses),
            "Products": len(analyses),
            "Evidence rows": len(evidence),
            "Warnings / review notes": sum(len(analysis.warnings) for analysis in analyses),
            "Anomalies": len(anomalies),
        }
        engineering_keys = [
            "dimensions",
            "parameters",
            "materials",
            "standards",
            "bom_items",
            "tables",
            "torque_requirements",
            "revisions",
            "drawing_views",
            "drawing_references",
            "components",
            "assemblies",
            "fasteners",
            "variants",
            "schematics",
            "identifications",
            "notes",
            "connections",
        ]
        for key in engineering_keys:
            totals[key.replace("_", " ").title()] = sum(getattr(getattr(analysis, "engineering", None), "counts", lambda: {})().get(key, 0) for analysis in analyses)
        for metric, value in totals.items():
            rows.append({"Metric": metric, "Value": value, "Meaning": self._summary_meaning(metric)})
        confidence_values = [source.confidence for analysis in analyses for source in self._engineering_sources(getattr(analysis, "engineering", None))]
        rows.append({"Metric": "High confidence", "Value": sum(value >= 0.75 for value in confidence_values), "Meaning": "Structured engineering objects extracted with confidence >= 0.75."})
        rows.append({"Metric": "Medium confidence", "Value": sum(0.55 <= value < 0.75 for value in confidence_values), "Meaning": "Structured engineering objects extracted with confidence between 0.55 and 0.75."})
        rows.append({"Metric": "Low confidence", "Value": sum(value < 0.55 for value in confidence_values), "Meaning": "Structured engineering objects that should be reviewed manually."})
        ocr_available = any(getattr(getattr(analysis, "engineering", None), "ocr_available", False) for analysis in analyses)
        ocr_used = any(getattr(getattr(analysis, "engineering", None), "ocr_used", False) for analysis in analyses)
        rows.append({"Metric": "OCR available", "Value": "YES" if ocr_available else "NO", "Meaning": "Whether Tesseract/pytesseract can be used for scanned drawings on this machine."})
        rows.append({"Metric": "OCR used", "Value": "YES" if ocr_used else "NO", "Meaning": "Whether any low-text page was actually analyzed through OCR in this run."})
        return pd.DataFrame(rows, columns=["Metric", "Value", "Meaning"])

    def _products_frame(self, analyses: list[Any], matches: list[Any]) -> pd.DataFrame:
        match_by_pdf = {match.source_pdf: match for match in matches}
        rows = []
        for analysis in analyses:
            match = match_by_pdf.get(analysis.source_pdf.name)
            rows.append(self._analysis_columns(analysis, match))
        return pd.DataFrame(rows)

    def _variants_frame(self, analyses: list[Any]) -> pd.DataFrame:
        rows = []
        for analysis in analyses:
            for variant in analysis.variants:
                rows.append({"Source PDF": analysis.source_pdf.name, "Variant Source": "base_parser", "Code": variant, "Variant Type": "detected_code", "Value": variant, "Page": "", "Confidence": ""})
            engineering = getattr(analysis, "engineering", None)
            for item in getattr(engineering, "variants", []) or []:
                rows.append({"Source PDF": analysis.source_pdf.name, "Variant Source": item.source.method, "Code": item.code, "Variant Type": item.variant_type, "Value": item.value, **self._source_dict(item.source)})
        return pd.DataFrame(rows, columns=["Source PDF", "Variant Source", "Code", "Variant Type", "Value", *self._source_headers()])

    def _bom_frame(self, analyses: list[Any]) -> pd.DataFrame:
        rows = []
        for analysis in analyses:
            for row in analysis.bom_rows:
                rows.append({"Source PDF": analysis.source_pdf.name, "Reference": "", "Part Number": row.get("ComponentPartNumber", ""), "Quantity": row.get("Quantity", ""), "Description": row.get("Description", ""), "Material": "", "Standard": row.get("Specification", ""), "Extraction Method": "base_bom_parser", "Page": "", "Confidence": "", "Evidence": row.get("EvidenceText", ""), "BBox X0": "", "BBox Y0": "", "BBox X1": "", "BBox Y1": ""})
            engineering = getattr(analysis, "engineering", None)
            for item in getattr(engineering, "bom_items", []) or []:
                rows.append({"Source PDF": analysis.source_pdf.name, "Reference": item.reference, "Part Number": item.part_number, "Quantity": item.quantity, "Description": item.description, "Material": item.material, "Standard": item.standard, **self._source_dict(item.source)})
        return pd.DataFrame(rows, columns=["Source PDF", "Reference", "Part Number", "Quantity", "Description", "Material", "Standard", *self._source_headers()])

    def _components_frame(self, analyses: list[Any]) -> pd.DataFrame:
        rows = []
        for analysis in analyses:
            engineering = getattr(analysis, "engineering", None)
            for item in getattr(engineering, "components", []) or []:
                rows.append({"Source PDF": analysis.source_pdf.name, "Reference": item.reference, "Part Number": item.part_number, "Description": item.description, "Quantity": item.quantity, "Material": item.material, **self._source_dict(item.source)})
        return pd.DataFrame(rows, columns=["Source PDF", "Reference", "Part Number", "Description", "Quantity", "Material", *self._source_headers()])

    def _dimensions_frame(self, analyses: list[Any]) -> pd.DataFrame:
        rows = []
        for analysis in analyses:
            engineering = getattr(analysis, "engineering", None)
            for item in getattr(engineering, "dimensions", []) or []:
                rows.append({"Source PDF": analysis.source_pdf.name, "Dimension Type": item.dimension_type, "Value": item.value, "Nominal": item.nominal_value, "Min": item.value_min, "Max": item.value_max, "Upper Tolerance": item.upper_tolerance, "Lower Tolerance": item.lower_tolerance, "Unit": item.unit, "Thread": item.thread_designation, **self._source_dict(item.source)})
        return pd.DataFrame(rows, columns=["Source PDF", "Dimension Type", "Value", "Nominal", "Min", "Max", "Upper Tolerance", "Lower Tolerance", "Unit", "Thread", *self._source_headers()])

    def _parameters_frame(self, analyses: list[Any]) -> pd.DataFrame:
        rows = []
        for analysis in analyses:
            engineering = getattr(analysis, "engineering", None)
            for item in getattr(engineering, "parameters", []) or []:
                rows.append({"Source PDF": analysis.source_pdf.name, "Parameter": item.name, "Value": item.value, "Min": item.value_min, "Max": item.value_max, "Unit": item.unit, "Tolerance": item.tolerance, "Qualifier": item.qualifier, **self._source_dict(item.source)})
        return pd.DataFrame(rows, columns=["Source PDF", "Parameter", "Value", "Min", "Max", "Unit", "Tolerance", "Qualifier", *self._source_headers()])

    def _materials_frame(self, analyses: list[Any]) -> pd.DataFrame:
        rows = []
        for analysis in analyses:
            engineering = getattr(analysis, "engineering", None)
            for item in getattr(engineering, "materials", []) or []:
                rows.append({"Source PDF": analysis.source_pdf.name, "Material": item.material, "Grade": item.grade, "Standard": item.standard, "Finish": item.finish, **self._source_dict(item.source)})
        return pd.DataFrame(rows, columns=["Source PDF", "Material", "Grade", "Standard", "Finish", *self._source_headers()])

    def _standards_frame(self, analyses: list[Any]) -> pd.DataFrame:
        rows = []
        for analysis in analyses:
            engineering = getattr(analysis, "engineering", None)
            for item in getattr(engineering, "standards", []) or []:
                rows.append({"Source PDF": analysis.source_pdf.name, "Standard": item.standard, "Applies To": item.applies_to, **self._source_dict(item.source)})
        return pd.DataFrame(rows, columns=["Source PDF", "Standard", "Applies To", *self._source_headers()])

    def _torque_fasteners_frame(self, analyses: list[Any]) -> pd.DataFrame:
        rows = []
        for analysis in analyses:
            engineering = getattr(analysis, "engineering", None)
            for item in getattr(engineering, "torque_requirements", []) or []:
                rows.append({"Source PDF": analysis.source_pdf.name, "Record Type": "torque", "Reference": item.reference, "Fastener Type": "", "Thread": item.thread, "Quantity": item.quantity, "Torque": item.torque, "Unit": item.unit, "Safety Class": item.safety_class, "Standard": "", "Instruction": item.instruction, **self._source_dict(item.source)})
            for item in getattr(engineering, "fasteners", []) or []:
                rows.append({"Source PDF": analysis.source_pdf.name, "Record Type": "fastener", "Reference": item.reference, "Fastener Type": item.fastener_type, "Thread": item.thread, "Quantity": item.quantity, "Torque": "", "Unit": "", "Safety Class": "", "Standard": item.standard, "Instruction": item.raw_text, **self._source_dict(item.source)})
        return pd.DataFrame(rows, columns=["Source PDF", "Record Type", "Reference", "Fastener Type", "Thread", "Quantity", "Torque", "Unit", "Safety Class", "Standard", "Instruction", *self._source_headers()])

    def _assemblies_frame(self, analyses: list[Any]) -> pd.DataFrame:
        rows = []
        for analysis in analyses:
            engineering = getattr(analysis, "engineering", None)
            for item in getattr(engineering, "assemblies", []) or []:
                rows.append({"Source PDF": analysis.source_pdf.name, "Assembly Name": item.name, "Assembly Type": item.assembly_type, "Component Refs": "; ".join(item.component_refs), **self._source_dict(item.source)})
        return pd.DataFrame(rows, columns=["Source PDF", "Assembly Name", "Assembly Type", "Component Refs", *self._source_headers()])

    def _views_frame(self, analyses: list[Any]) -> pd.DataFrame:
        rows = []
        for analysis in analyses:
            engineering = getattr(analysis, "engineering", None)
            for item in getattr(engineering, "drawing_views", []) or []:
                rows.append({"Source PDF": analysis.source_pdf.name, "Label": item.label, "View Type": item.view_type, "Scale": item.scale, **self._source_dict(item.source)})
        return pd.DataFrame(rows, columns=["Source PDF", "Label", "View Type", "Scale", *self._source_headers()])

    def _references_frame(self, analyses: list[Any]) -> pd.DataFrame:
        rows = []
        for analysis in analyses:
            engineering = getattr(analysis, "engineering", None)
            for item in getattr(engineering, "drawing_references", []) or []:
                rows.append({"Source PDF": analysis.source_pdf.name, "Reference": item.reference, "Related Part Number": item.related_part_number, "Relationship Confidence": item.relationship_confidence, **self._source_dict(item.source)})
        return pd.DataFrame(rows, columns=["Source PDF", "Reference", "Related Part Number", "Relationship Confidence", *self._source_headers()])

    def _connections_frame(self, analyses: list[Any]) -> pd.DataFrame:
        rows = []
        for analysis in analyses:
            engineering = getattr(analysis, "engineering", None)
            for item in getattr(engineering, "connections", []) or []:
                rows.append({"Source PDF": analysis.source_pdf.name, "Function": item.function, "Thread Type": item.thread_type, "Thread Size": item.thread_size, "Pitch": item.pitch, "Standard": item.standard, **self._source_dict(item.source)})
        return pd.DataFrame(rows, columns=["Source PDF", "Function", "Thread Type", "Thread Size", "Pitch", "Standard", *self._source_headers()])

    def _schematics_frame(self, analyses: list[Any]) -> pd.DataFrame:
        rows = []
        for analysis in analyses:
            engineering = getattr(analysis, "engineering", None)
            for item in getattr(engineering, "schematics", []) or []:
                rows.append({"Source PDF": analysis.source_pdf.name, "Label": item.label, "Function": item.function, "Connection": item.connection, **self._source_dict(item.source)})
        return pd.DataFrame(rows, columns=["Source PDF", "Label", "Function", "Connection", *self._source_headers()])

    def _notes_frame(self, analyses: list[Any]) -> pd.DataFrame:
        rows = []
        for analysis in analyses:
            engineering = getattr(analysis, "engineering", None)
            for item in getattr(engineering, "notes", []) or []:
                rows.append({"Source PDF": analysis.source_pdf.name, "Category": item.category, "Text": item.text, **self._source_dict(item.source)})
        return pd.DataFrame(rows, columns=["Source PDF", "Category", "Text", *self._source_headers()])

    def _revisions_frame(self, analyses: list[Any]) -> pd.DataFrame:
        rows = []
        for analysis in analyses:
            engineering = getattr(analysis, "engineering", None)
            for item in getattr(engineering, "revisions", []) or []:
                rows.append({"Source PDF": analysis.source_pdf.name, "Revision": item.revision, "Date": item.date, "Change Type": item.change_type, "Affected Reference": item.affected_reference, "Description": item.description, **self._source_dict(item.source)})
        return pd.DataFrame(rows, columns=["Source PDF", "Revision", "Date", "Change Type", "Affected Reference", "Description", *self._source_headers()])

    def _identification_frame(self, analyses: list[Any]) -> pd.DataFrame:
        rows = []
        for analysis in analyses:
            rows.append({"Source PDF": analysis.source_pdf.name, "Identifier Type": "drawing_number", "Value": analysis.fields.get("Drawing number", ""), "Evidence": "", "Page": "", "Extraction Method": "base_parser", "Confidence": "", "BBox X0": "", "BBox Y0": "", "BBox X1": "", "BBox Y1": ""})
            engineering = getattr(analysis, "engineering", None)
            for item in getattr(engineering, "identifications", []) or []:
                rows.append({"Source PDF": analysis.source_pdf.name, "Identifier Type": item.identifier_type, "Value": item.value, **self._source_dict(item.source)})
        return pd.DataFrame(rows, columns=["Source PDF", "Identifier Type", "Value", *self._source_headers()])

    def _sources_frame(self, analyses: list[Any]) -> pd.DataFrame:
        rows = []
        for analysis in analyses:
            for ev in analysis.evidence:
                rows.append({"Source PDF": ev.source_pdf, "Entity": ev.field, "Value": ev.value, "Status": ev.status, "Page": ev.page, "Extraction Method": ev.extraction_method, "Confidence": ev.confidence, "Evidence": ev.evidence_text, "BBox X0": "", "BBox Y0": "", "BBox X1": "", "BBox Y1": ""})
            engineering = getattr(analysis, "engineering", None)
            for source in self._engineering_sources(engineering):
                rows.append({"Source PDF": source.source_pdf, "Entity": source.region_type, "Value": source.raw_text[:180], "Status": "ACCEPTED" if source.confidence >= 0.6 else "REVIEW", **self._source_dict(source)})
        return pd.DataFrame(rows, columns=["Source PDF", "Entity", "Value", "Status", *self._source_headers()])

    def _quality_frame(self, analyses: list[Any], anomalies: list[dict[str, Any]]) -> pd.DataFrame:
        rows = []
        for analysis in analyses:
            for warning in analysis.warnings:
                rows.append({"Entity": "Validation", "Type": "warning", "Document": analysis.source_pdf.name, "Page": "", "Value": "", "Confidence": "", "Method": "VALIDATION", "Warning": warning, "Review Required": "YES"})
            engineering = getattr(analysis, "engineering", None)
            for source in self._engineering_sources(engineering):
                if source.confidence < 0.55:
                    rows.append({"Entity": source.region_type, "Type": "low_confidence", "Document": analysis.source_pdf.name, "Page": source.page, "Value": source.raw_text[:180], "Confidence": source.confidence, "Method": source.method, "Warning": f"Confidence {source.confidence:.2f}", "Review Required": "YES"})
        for anomaly in anomalies:
            rows.append({"Entity": "Catalogue anomaly", "Type": anomaly.get("severity", "Anomaly"), "Document": anomaly.get("Source PDF", ""), "Page": "", "Value": "", "Confidence": "", "Method": "ANOMALY", "Warning": anomaly.get("message", str(anomaly)), "Review Required": "YES"})
        return pd.DataFrame(rows, columns=["Entity", "Type", "Document", "Page", "Value", "Confidence", "Method", "Warning", "Review Required"])

    def _raw_data_frame(self, analyses: list[Any]) -> pd.DataFrame:
        rows = []
        for analysis in analyses:
            engineering = getattr(analysis, "engineering", None)
            for item in getattr(engineering, "raw_data", []) or []:
                bbox = item.bbox or ("", "", "", "")
                width = round(bbox[2] - bbox[0], 2) if all(isinstance(value, (int, float)) for value in bbox) else ""
                height = round(bbox[3] - bbox[1], 2) if all(isinstance(value, (int, float)) for value in bbox) else ""
                rows.append({"Source PDF": analysis.source_pdf.name, "Page": item.page, "Region Type": item.region_type, "Method": item.method, "Confidence": item.confidence, "Text": item.text, "BBox X0": bbox[0], "BBox Y0": bbox[1], "BBox X1": bbox[2], "BBox Y1": bbox[3], "BBox Width": width, "BBox Height": height})
        return pd.DataFrame(rows, columns=["Source PDF", "Page", "Region Type", "Method", "Confidence", "Text", "BBox X0", "BBox Y0", "BBox X1", "BBox Y1", "BBox Width", "BBox Height"])

    def _engineering_sources(self, engineering: Any) -> list[Any]:
        if not engineering:
            return []
        sources = []
        for attr in ("dimensions", "parameters", "materials", "standards", "bom_items", "tables", "torque_requirements", "revisions", "drawing_views", "drawing_references", "components", "assemblies", "fasteners", "variants", "schematics", "identifications", "notes", "connections"):
            for item in getattr(engineering, attr, []) or []:
                source = getattr(item, "source", None)
                if source:
                    sources.append(source)
        return sources

    def _source_headers(self) -> list[str]:
        return ["Page", "Extraction Method", "Confidence", "Evidence", "BBox X0", "BBox Y0", "BBox X1", "BBox Y1", "BBox Width", "BBox Height"]

    def _source_dict(self, source: Any) -> dict[str, Any]:
        bbox = source.bbox or ("", "", "", "")
        width = round(bbox[2] - bbox[0], 2) if all(isinstance(value, (int, float)) for value in bbox) else ""
        height = round(bbox[3] - bbox[1], 2) if all(isinstance(value, (int, float)) for value in bbox) else ""
        return {
            "Page": source.page,
            "Extraction Method": source.method,
            "Confidence": source.confidence,
            "Evidence": source.raw_text,
            "BBox X0": bbox[0],
            "BBox Y0": bbox[1],
            "BBox X1": bbox[2],
            "BBox Y1": bbox[3],
            "BBox Width": width,
            "BBox Height": height,
        }

    def _summary_meaning(self, metric: str) -> str:
        meanings = {
            "PDF drawings analyzed": "Input technical drawings processed in this run.",
            "Pages processed": "Total PDF pages parsed through the analyzer.",
            "Products": "Product/drawing records written to the catalogue view.",
            "Evidence rows": "Traceability rows connecting workbook values to source text.",
            "Warnings / review notes": "Items that should be manually checked before client release.",
            "Anomalies": "Detected catalogue consistency or rule issues.",
            "Dimensions": "Extracted quotas, holes, radii, angles, threads, and toleranced values.",
            "Parameters": "Pressure, temperature, electrical, performance, weight, and similar technical values.",
            "Bom Items": "Parts list rows detected from drawing text or PDF table regions.",
            "Drawing References": "Balloon/item references detected in drawings.",
        }
        return meanings.get(metric, "Structured engineering data extracted from the PDFs.")

    def _write_summary_sheet(self, ws: Any, frame: pd.DataFrame) -> None:
        self._write_dataframe_sheet(ws, "SUMMARY", frame)
        if frame.empty or "Value" not in frame.columns:
            return
        max_row = ws.max_row
        chart = BarChart()
        chart.title = "Engineering extraction coverage"
        chart.y_axis.title = "Count"
        chart.x_axis.title = "Entity"
        data = Reference(ws, min_col=2, min_row=4, max_row=min(max_row, 22))
        cats = Reference(ws, min_col=1, min_row=5, max_row=min(max_row, 22))
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.height = 8
        chart.width = 18
        ws.add_chart(chart, "E4")

    def _write_dataframe_sheet(self, ws: Any, title: str, frame: pd.DataFrame) -> None:
        frame = frame.copy()
        frame = frame.fillna("")
        ws.sheet_view.showGridLines = False
        max_col = max(1, len(frame.columns))
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
        title_cell = ws.cell(row=1, column=1, value=title)
        title_cell.fill = PatternFill("solid", fgColor="111827")
        title_cell.font = Font(color="FFFFFF", bold=True, size=14)
        title_cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[1].height = 26
        ws.cell(row=2, column=1, value=f"{len(frame)} extracted rows. All values keep page/source provenance where available.")
        ws.cell(row=2, column=1).font = Font(color="64748B", italic=True)
        header_row = 4
        for col_idx, column in enumerate(frame.columns, start=1):
            cell = ws.cell(row=header_row, column=col_idx, value=column)
            cell.fill = PatternFill("solid", fgColor="E8EEF7")
            cell.font = Font(color="111827", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=Side(style="thin", color="CBD5E1"))
        for row_idx, record in enumerate(frame.to_dict("records"), start=header_row + 1):
            for col_idx, column in enumerate(frame.columns, start=1):
                value = record.get(column, "")
                if isinstance(value, (list, tuple, set)):
                    value = "; ".join(str(item) for item in value)
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if row_idx % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F8FAFC")
        ws.freeze_panes = "A5"
        if not frame.empty:
            table_ref = f"A{header_row}:{get_column_letter(max_col)}{ws.max_row}"
            table_name = re.sub(r"[^A-Za-z0-9]", "", title.title())[:24] or "Data"
            table = Table(displayName=f"{table_name}Table", ref=table_ref)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
            ws.add_table(table)
        else:
            ws.auto_filter.ref = f"A{header_row}:{get_column_letter(max_col)}{header_row}"
        for col_idx, column in enumerate(frame.columns, start=1):
            series = [str(column), *(str(value) for value in frame[column].head(150).tolist())] if column in frame.columns else [str(column)]
            width = min(48, max(10, max(len(value) for value in series[:151]) + 2))
            if column in {"Evidence", "Text", "Description", "Instruction", "Issue", "Action", "Meaning", "PDF warnings", "Analysis note"}:
                width = 42
            ws.column_dimensions[get_column_letter(col_idx)].width = width
        for row_idx in range(header_row + 1, min(ws.max_row, header_row + 80) + 1):
            ws.row_dimensions[row_idx].height = 32
        if "Confidence" in frame.columns and not frame.empty:
            col_letter = get_column_letter(list(frame.columns).index("Confidence") + 1)
            ws.conditional_formatting.add(
                f"{col_letter}{header_row + 1}:{col_letter}{ws.max_row}",
                ColorScaleRule(start_type="num", start_value=0, start_color="FCA5A5", mid_type="num", mid_value=0.65, mid_color="FDE68A", end_type="num", end_value=1, end_color="86EFAC"),
            )
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    def _build_technical_characteristics_frame(self, analyses: list[Any]) -> pd.DataFrame:
        fields = [
            "Diameter",
            "LED",
            "Pressure",
            "Mounting",
            "Fitting",
            "Accuracy class",
            "Earth lug",
            "Outlet connection",
            "Envelope dimensions",
            "Mounting holes",
            "Key slot",
            "Working pressure",
            "Working temperature",
            "Startup temperature",
            "Duty cycle",
            "Starts per hour",
            "Weight",
            "Configuration",
            "Compressor detail",
            "Protection degree",
            "Electrical rating",
            "Illumination detail",
            "Compressed air quality",
            "Case material",
            "Frame ring",
            "Dial",
            "Pointer",
            "Pointer system",
            "Color coding",
            "Process connection",
            "Lens",
            "Restrictor screw",
            "Safety blow-out",
            "Standards",
            "Material finishes",
        ]
        rows: list[dict[str, Any]] = []
        for analysis in analyses:
            evidence_lookup = self._technical_evidence_lookup(analysis)
            for field in fields:
                value = analysis.fields.get(field, "")
                if not value or value in {"UNKNOWN", "NOT_FOUND", "NEEDS REVIEW"}:
                    continue
                evidence = evidence_lookup.get((field, value), {})
                rows.append(
                    {
                        "Source PDF": analysis.source_pdf.name,
                        "Drawing number": analysis.fields.get("Drawing number", ""),
                        "Product Name": analysis.fields.get("Product Name", ""),
                        "Product Type": analysis.fields.get("Product Type", ""),
                        "Category": self._technical_category(field),
                        "Characteristic": field,
                        "Mapped value": value,
                        "Confidence": evidence.get("confidence", ""),
                        "Evidence": evidence.get("evidence_text", ""),
                    }
                )
            for key, value in analysis.fields.items():
                if not key.startswith("Technical characteristic|") or not value:
                    continue
                _, category, name, _idx = key.split("|", 3)
                evidence = evidence_lookup.get((f"Technical characteristic: {category} / {name}", value), {})
                rows.append(
                    {
                        "Source PDF": analysis.source_pdf.name,
                        "Drawing number": analysis.fields.get("Drawing number", ""),
                        "Product Name": analysis.fields.get("Product Name", ""),
                        "Product Type": analysis.fields.get("Product Type", ""),
                        "Category": category,
                        "Characteristic": name,
                        "Mapped value": value,
                        "Confidence": evidence.get("confidence", ""),
                        "Evidence": evidence.get("evidence_text", ""),
                    }
                )
        columns = ["Source PDF", "Drawing number", "Product Name", "Product Type", "Category", "Characteristic", "Mapped value", "Confidence", "Evidence"]
        frame = pd.DataFrame(rows, columns=columns)
        if frame.empty:
            return frame
        return frame.drop_duplicates(subset=["Source PDF", "Category", "Characteristic", "Mapped value"]).sort_values(["Source PDF", "Category", "Characteristic", "Mapped value"])

    def _build_technical_summary_frame(self, technical_rows: pd.DataFrame) -> pd.DataFrame:
        columns = ["Source PDF", "Category", "Extracted values", "Average confidence", "Examples"]
        if technical_rows.empty:
            return pd.DataFrame(columns=columns)
        summary = (
            technical_rows.assign(Confidence=pd.to_numeric(technical_rows["Confidence"], errors="coerce"))
            .groupby(["Source PDF", "Category"], dropna=False)
            .agg(
                **{
                    "Extracted values": ("Mapped value", "nunique"),
                    "Average confidence": ("Confidence", "mean"),
                    "Examples": ("Mapped value", lambda values: "; ".join(str(value) for value in list(values)[:5])),
                }
            )
            .reset_index()
        )
        summary["Average confidence"] = summary["Average confidence"].round(2).fillna("")
        return summary[columns].sort_values(["Source PDF", "Category"])

    def _build_client_catalogue(self, generated: pd.DataFrame, analyses: list[Any]) -> pd.DataFrame:
        rows: list[dict[str, Any]] = []
        for row, analysis in zip(generated.to_dict("records"), analyses):
            enriched = dict(row)
            engineering_summary = self._engineering_summary(analysis)
            enriched["Technical attribute 5"] = self._join_known_limited(
                analysis.fields.get("Accuracy class"),
                analysis.fields.get("Protection degree"),
                max_items=4,
            )
            enriched["Technical attribute 6"] = self._join_known_limited(
                analysis.fields.get("Electrical rating"),
                analysis.fields.get("Illumination detail"),
                max_items=4,
            )
            enriched["Technical attribute 7"] = self._join_known_limited(
                analysis.fields.get("Fitting"),
                analysis.fields.get("Process connection"),
                max_items=4,
            )
            enriched["Technical attribute 8"] = self._join_known_limited(
                analysis.fields.get("Working temperature") or analysis.fields.get("Temperature"),
                analysis.fields.get("Compressed air quality"),
                engineering_summary.get("parameters"),
                max_items=4,
            )
            enriched["Technical attribute 9"] = self._join_known_limited(
                analysis.fields.get("Case material"),
                analysis.fields.get("Frame ring"),
                analysis.fields.get("Dial"),
                analysis.fields.get("Material finishes"),
                max_items=5,
            )
            enriched["Technical attribute 10"] = self._join_known_limited(
                engineering_summary.get("dimensions"),
                analysis.fields.get("Pointer"),
                analysis.fields.get("Pointer system"),
                analysis.fields.get("Color coding"),
                max_items=4,
            )
            enriched["Technical attribute 11"] = self._join_known_limited(
                analysis.fields.get("Lens"),
                analysis.fields.get("Restrictor screw"),
                analysis.fields.get("Safety blow-out"),
                engineering_summary.get("torque_fasteners"),
                engineering_summary.get("bom"),
                max_items=4,
            )
            enriched["Technical attribute 12"] = self._join_known_limited(
                analysis.fields.get("Standards"),
                engineering_summary.get("revisions_views"),
                max_items=8,
            )
            rows.append(enriched)

        columns = list(generated.columns)
        insert_at = columns.index("Technical attribute 4") + 1 if "Technical attribute 4" in columns else len(columns)
        extra_columns = [f"Technical attribute {number}" for number in range(5, 13)]
        for offset, column in enumerate(extra_columns):
            columns.insert(insert_at + offset, column)
        return pd.DataFrame(rows, columns=columns)

    def _engineering_summary(self, analysis: Any) -> dict[str, str]:
        engineering = getattr(analysis, "engineering", None)
        if not engineering:
            return {}
        dimensions = []
        for dimension in engineering.dimensions[:8]:
            if dimension.dimension_type == "thread":
                dimensions.append(dimension.thread_designation)
            elif dimension.dimension_type in {"diameter", "radius", "angle", "slot"}:
                dimensions.append(f"{dimension.dimension_type}: {dimension.value}")
            elif dimension.upper_tolerance is not None:
                dimensions.append(f"{dimension.value} {dimension.unit}".strip())
        parameters = []
        for parameter in engineering.parameters[:8]:
            label = parameter.name
            value = f"{parameter.value} {parameter.unit}".strip()
            if parameter.tolerance:
                value = f"{value} {parameter.tolerance}"
            parameters.append(f"{label}: {value}")
        torque = [f"{item.reference + ' ' if item.reference else ''}{item.thread + ' ' if item.thread else ''}{item.torque} {item.unit}".strip() for item in engineering.torque_requirements[:5]]
        fasteners = [connection.raw_text for connection in engineering.connections[:5] if re.search(r"screw|bolt|nut|washer|helicoil|thread", connection.raw_text, re.I)]
        bom = []
        if engineering.bom_items:
            bom.append(f"BOM items: {len(engineering.bom_items)}")
            bom.extend(item.part_number for item in engineering.bom_items[:4] if item.part_number)
        revisions = [f"REV {item.revision}: {item.change_type}" for item in engineering.revisions[:5]]
        views = [f"{item.view_type}: {item.label}" for item in engineering.drawing_views[:5]]
        return {
            "dimensions": "; ".join(dict.fromkeys(dimensions)),
            "parameters": "; ".join(dict.fromkeys(parameters)),
            "torque_fasteners": "; ".join(dict.fromkeys([*torque, *fasteners])),
            "bom": "; ".join(dict.fromkeys(bom)),
            "revisions_views": "; ".join(dict.fromkeys([*revisions, *views])),
        }

    def _technical_evidence_lookup(self, analysis: Any) -> dict[tuple[str, str], dict[str, Any]]:
        lookup: dict[tuple[str, str], dict[str, Any]] = {}
        for item in analysis.evidence:
            lookup[(item.field, item.value)] = {
                "confidence": item.confidence,
                "evidence_text": item.evidence_text,
            }
        return lookup

    def _technical_category(self, field: str) -> str:
        categories = {
            "Diameter": "dimension",
            "LED": "illumination",
            "Pressure": "pressure",
            "Mounting": "mounting",
            "Fitting": "connection",
            "Accuracy class": "standard",
            "Earth lug": "connection",
            "Outlet connection": "connection",
            "Envelope dimensions": "dimension",
            "Mounting holes": "connection",
            "Key slot": "dimension",
            "Working pressure": "pressure",
            "Working temperature": "temperature",
            "Startup temperature": "temperature",
            "Duty cycle": "performance",
            "Starts per hour": "performance",
            "Weight": "weight",
            "Configuration": "configuration",
            "Compressor detail": "configuration",
            "Protection degree": "protection",
            "Electrical rating": "electrical",
            "Illumination detail": "illumination",
            "Compressed air quality": "media",
            "Case material": "material",
            "Frame ring": "material",
            "Dial": "material",
            "Pointer": "material",
            "Pointer system": "component",
            "Color coding": "component",
            "Process connection": "connection",
            "Lens": "component",
            "Restrictor screw": "component",
            "Safety blow-out": "component",
            "Standards": "standard",
            "Material finishes": "material",
        }
        return categories.get(field, "technical")

    def _write_template_catalogue_sheet(self, path: Path, generated: pd.DataFrame) -> None:
        wb = load_workbook(self.settings.catalogue_path)
        ws = wb[wb.sheetnames[0]]
        header_row = 4
        first_data_row = 5
        first_catalogue_col = 2
        original_headers = [ws.cell(header_row, column).value for column in range(first_catalogue_col, ws.max_column + 1)]
        headers = list(generated.columns)
        added_headers = [header for header in headers if header not in original_headers]
        if added_headers and "Technical attribute 4" in original_headers:
            insert_at = first_catalogue_col + original_headers.index("Technical attribute 4") + 1
            source_col = insert_at - 1
            ws.insert_cols(insert_at, amount=len(added_headers))
            for column in range(insert_at, insert_at + len(added_headers)):
                source_letter = ws.cell(header_row, source_col).column_letter
                target_letter = ws.cell(header_row, column).column_letter
                ws.column_dimensions[target_letter].width = ws.column_dimensions[source_letter].width
                for row in range(1, first_data_row + 1):
                    target = ws.cell(row, column)
                    source = ws.cell(row, source_col)
                    target._style = copy(source._style)
                    target.number_format = source.number_format
                    target.alignment = copy(source.alignment)
                    target.protection = copy(source.protection)
                    target.fill = copy(source.fill)
                    target.font = copy(source.font)
                    target.border = copy(source.border)
        header_style_by_column = {
            column: {
                "style": copy(ws.cell(header_row, column)._style),
                "number_format": ws.cell(header_row, column).number_format,
                "alignment": copy(ws.cell(header_row, column).alignment),
                "protection": copy(ws.cell(header_row, column).protection),
                "fill": copy(ws.cell(header_row, column).fill),
                "font": copy(ws.cell(header_row, column).font),
                "border": copy(ws.cell(header_row, column).border),
            }
            for column in range(first_catalogue_col, first_catalogue_col + len(headers))
        }
        data_style_source_row = first_data_row if ws.max_row >= first_data_row else header_row
        data_style_by_column = {
            column: {
                "style": copy(ws.cell(data_style_source_row, column)._style),
                "number_format": ws.cell(data_style_source_row, column).number_format,
                "alignment": copy(ws.cell(data_style_source_row, column).alignment),
                "protection": copy(ws.cell(data_style_source_row, column).protection),
                "fill": copy(ws.cell(data_style_source_row, column).fill),
                "font": copy(ws.cell(data_style_source_row, column).font),
                "border": copy(ws.cell(data_style_source_row, column).border),
            }
            for column in range(first_catalogue_col, first_catalogue_col + len(headers))
        }
        for column_offset, header in enumerate(headers, start=first_catalogue_col):
            cell = ws.cell(header_row, column_offset)
            cell.value = header
            if column_offset in header_style_by_column:
                source = header_style_by_column[column_offset]
                cell._style = copy(source["style"])
                cell.number_format = source["number_format"]
                cell.alignment = copy(source["alignment"])
                cell.protection = copy(source["protection"])
                cell.fill = copy(source["fill"])
                cell.font = copy(source["font"])
                cell.border = copy(source["border"])
        if ws.max_row >= first_data_row:
            ws.delete_rows(first_data_row, ws.max_row - first_data_row + 1)
        for row_offset, record in enumerate(generated.to_dict("records"), start=first_data_row):
            for column_offset, header in enumerate(headers, start=first_catalogue_col):
                cell = ws.cell(row_offset, column_offset)
                cell.value = record.get(header, "")
                source = data_style_by_column[column_offset]
                cell._style = copy(source["style"])
                cell.number_format = source["number_format"]
                cell.alignment = copy(source["alignment"])
                cell.protection = copy(source["protection"])
                cell.fill = copy(source["fill"])
                cell.font = copy(source["font"])
                cell.border = copy(source["border"])
        ws.freeze_panes = "B5"
        wb.save(path)

    def _map_technical_attributes(self, row: dict[str, Any], analysis: Any, prefer_existing: bool = False) -> None:
        choose = self._existing_or_known if prefer_existing else self._known_or_existing
        product_name = analysis.fields.get("Product Name", "")
        if product_name == "D - MANOMETERS":
            row["Technical attribute 1"] = choose(row.get("Technical attribute 1"), analysis.fields.get("Diameter")) if prefer_existing else choose(analysis.fields.get("Diameter"), row.get("Technical attribute 1"))
            row["Technical attribute 2"] = choose(row.get("Technical attribute 2"), analysis.fields.get("LED")) if prefer_existing else choose(analysis.fields.get("LED"), row.get("Technical attribute 2"))
            row["Technical attribute 3"] = choose(row.get("Technical attribute 3"), analysis.fields.get("Pressure")) if prefer_existing else choose(analysis.fields.get("Pressure"), row.get("Technical attribute 3"))
            row["Technical attribute 4"] = choose(row.get("Technical attribute 4"), analysis.fields.get("Mounting")) if prefer_existing else choose(analysis.fields.get("Mounting"), row.get("Technical attribute 4"))
            return
        if product_name == "A-BURAN COMPRESSOR":
            row["Technical attribute 1"] = self._first_known(analysis.fields.get("Outlet connection"), analysis.fields.get("Compressor detail"), row.get("Technical attribute 1"))
            row["Technical attribute 2"] = self._first_known(analysis.fields.get("Envelope dimensions"), row.get("Technical attribute 2"))
            row["Technical attribute 3"] = self._join_known(
                analysis.fields.get("Weight"),
                analysis.fields.get("Working pressure"),
                analysis.fields.get("Working temperature"),
                analysis.fields.get("Duty cycle"),
                fallback=row.get("Technical attribute 3"),
            )
            row["Technical attribute 4"] = self._join_known(analysis.fields.get("Mounting holes"), analysis.fields.get("Key slot"), fallback=row.get("Technical attribute 4"))
            return
        row["Technical attribute 1"] = choose(row.get("Technical attribute 1"), analysis.fields.get("Diameter")) if prefer_existing else choose(analysis.fields.get("Diameter"), row.get("Technical attribute 1"))
        row["Technical attribute 2"] = choose(row.get("Technical attribute 2"), analysis.fields.get("Drain")) if prefer_existing else choose(analysis.fields.get("Drain"), row.get("Technical attribute 2"))
        row["Technical attribute 3"] = choose(row.get("Technical attribute 3"), analysis.fields.get("Contact")) if prefer_existing else choose(analysis.fields.get("Contact"), row.get("Technical attribute 3"))
        row["Technical attribute 4"] = choose(row.get("Technical attribute 4"), analysis.fields.get("Handle")) if prefer_existing else choose(analysis.fields.get("Handle"), row.get("Technical attribute 4"))

    def _focus_catalogue(self, catalogue: pd.DataFrame) -> pd.DataFrame:
        focus_value = self.settings.focus_product_name.upper()
        if "Product Name" not in catalogue.columns:
            return catalogue.iloc[0:0].copy()
        mask = catalogue["Product Name"].astype(str).str.upper().eq(focus_value)
        return catalogue.loc[mask].copy()

    def _build_demo_focus(self, catalogue: pd.DataFrame, focus_catalogue: pd.DataFrame, expected_count: int | None = None) -> dict[str, Any]:
        technical = [c for c in catalogue.columns if c.startswith("Technical attribute")]
        preview_columns = [c for c in ["PartNumber", "Master PN", "Product Family", "Product Name", "Product Type", *technical] if c in focus_catalogue.columns]
        attribute_counts: dict[str, dict[str, int]] = {}
        for column in ["Product Type", *technical]:
            if column in focus_catalogue.columns:
                counts = focus_catalogue[column].replace("", "UNKNOWN").value_counts().head(12)
                attribute_counts[column] = {str(k): int(v) for k, v in counts.items()}
        found = int(len(focus_catalogue))
        expected = int(expected_count if expected_count is not None else found)
        return {
            "brief": [
                "Use the initial Excel catalogue as a schema/reference template only.",
                "Read the supplied PDF drawings and produce a new catalogue from those PDFs.",
                "Map extracted drawing characteristics into the catalogue-compatible Excel columns.",
                "Run requested 1st-level and 2nd-level diagrams on the PDF-generated catalogue.",
            ],
            "filter_column": "Source",
            "filter_value": "PDF-generated catalogue",
            "expected_codes": expected,
            "found_codes": found,
            "missing_vs_expected": max(0, expected - found),
            "technical_columns": technical,
            "attribute_counts": attribute_counts,
            "records": focus_catalogue[preview_columns].to_dict("records") if preview_columns else [],
        }

    def _primary_part(self, part_numbers: list[str], fallback: str) -> str:
        if "XX" in fallback:
            return fallback
        for part in part_numbers:
            if "XX" not in part and part not in {"TBA", "N/A"}:
                return part
        return fallback

    def _master_part(self, part_numbers: list[str], fallback: str) -> str:
        for part in part_numbers:
            if "XX" in part:
                return part
        if fallback.startswith("FT") and fallback.endswith("-100"):
            return fallback[:-3] + "XXX"
        return fallback

    def _matched_catalogue_base(self, catalogue: pd.DataFrame, match: Any) -> dict[str, Any]:
        if not match or not match.matched_part_number:
            return {column: "" for column in catalogue.columns}
        rows = catalogue[catalogue["PartNumber"] == match.matched_part_number]
        if rows.empty:
            return {column: "" for column in catalogue.columns}
        return rows.iloc[0].to_dict()

    def _known_or_existing(self, extracted: str | None, existing: str | None) -> str:
        if extracted and extracted not in {"UNKNOWN", "NOT_FOUND"}:
            return extracted
        return existing or extracted or "UNKNOWN"

    def _first_known(self, *values: Any) -> str:
        for value in values:
            if value and str(value) not in {"UNKNOWN", "NOT_FOUND", "NEEDS REVIEW"}:
                return str(value)
        return ""

    def _join_known(self, *values: Any, fallback: Any = "") -> str:
        known = self._unique_known_values(*values)
        return "; ".join(known) if known else (str(fallback) if fallback else "")

    def _join_known_limited(self, *values: Any, fallback: Any = "", max_items: int = 4) -> str:
        known = self._unique_known_values(*values)
        if known:
            return "; ".join(known[:max_items])
        return str(fallback) if fallback else ""

    def _unique_known_values(self, *values: Any) -> list[str]:
        known: list[str] = []
        seen: set[str] = set()
        for value in values:
            if not value or str(value) in {"UNKNOWN", "NOT_FOUND", "NEEDS REVIEW"}:
                continue
            for part in str(value).split(";"):
                cleaned = part.strip()
                if not cleaned or cleaned in {"UNKNOWN", "NOT_FOUND", "NEEDS REVIEW"}:
                    continue
                key = cleaned.lower().replace(" ", "").replace(",", ".")
                if key in seen:
                    continue
                seen.add(key)
                known.append(cleaned)
        return known

    def _existing_or_known(self, existing: str | None, extracted: str | None) -> str:
        if existing and existing not in {"UNKNOWN", "NOT_FOUND"}:
            return existing
        return self._known_or_existing(extracted, existing)

    def _build_evidence_frame(self, analyses: list[Any]) -> pd.DataFrame:
        rows: list[dict[str, Any]] = []
        for analysis in analyses:
            for ev in analysis.evidence:
                rows.append(ev.__dict__)
            rows.extend(self._engineering_evidence_rows(analysis))
            for field, value in analysis.fields.items():
                if value in {"UNKNOWN", "NOT_FOUND", "NEEDS REVIEW"}:
                    rows.append(Evidence(field, value, analysis.source_pdf.name, 0, "Field not unambiguously available in deterministic extraction.", 0.25, "status_marker", "NEEDS REVIEW").__dict__)
        return pd.DataFrame(rows)

    def _engineering_evidence_rows(self, analysis: Any) -> list[dict[str, Any]]:
        engineering = getattr(analysis, "engineering", None)
        if not engineering:
            return []
        rows: list[dict[str, Any]] = []
        entity_specs = [
            ("Engineering dimension", engineering.dimensions, lambda item: item.value),
            ("Engineering parameter", engineering.parameters, lambda item: f"{item.name}: {item.value} {item.unit}".strip()),
            ("Engineering material", engineering.materials, lambda item: item.material),
            ("Engineering standard", engineering.standards, lambda item: item.standard),
            ("Engineering BOM item", engineering.bom_items, lambda item: item.part_number or item.description),
            ("Engineering table", engineering.tables, lambda item: f"{item.table_id}: {item.region_type}"),
            ("Engineering torque", engineering.torque_requirements, lambda item: f"{item.reference} {item.thread} {item.torque} {item.unit}".strip()),
            ("Engineering revision", engineering.revisions, lambda item: f"{item.revision}: {item.change_type}"),
            ("Engineering drawing view", engineering.drawing_views, lambda item: f"{item.view_type}: {item.label}"),
            ("Engineering drawing reference", engineering.drawing_references, lambda item: item.reference),
            ("Engineering component", engineering.components, lambda item: item.part_number or item.description),
            ("Engineering assembly", engineering.assemblies, lambda item: item.name),
            ("Engineering fastener", engineering.fasteners, lambda item: f"{item.fastener_type} {item.thread}".strip()),
            ("Engineering variant", engineering.variants, lambda item: item.code or item.value),
            ("Engineering schematic", engineering.schematics, lambda item: f"{item.function}: {item.label}"),
            ("Engineering identification", engineering.identifications, lambda item: f"{item.identifier_type}: {item.value}"),
            ("Engineering note", engineering.notes, lambda item: item.text),
            ("Engineering connection", engineering.connections, lambda item: item.raw_text),
        ]
        for field, items, value_fn in entity_specs:
            for item in items:
                source = item.source
                rows.append(
                    Evidence(
                        field,
                        value_fn(item),
                        source.source_pdf,
                        source.page,
                        source.raw_text,
                        source.confidence,
                        source.method,
                    ).__dict__
                )
        return rows

    def _write_json(self, path: Path, analyses: list[Any], matches: list[Any], rule_results: dict[str, Any], anomalies: list[dict[str, Any]], clusters: dict[str, Any], similarity: list[dict[str, Any]], bom_summary: dict[str, Any], demo_focus: dict[str, Any]) -> None:
        payload = {
            "demo_focus": demo_focus,
            "pdfs": [
                {
                    "source_pdf": a.source_pdf.name,
                    "page_count": a.page_count,
                    "fields": a.fields,
                    "part_numbers": a.part_numbers,
                    "variants": a.variants,
                    "engineering": a.engineering.to_dict() if getattr(a, "engineering", None) else {},
                    "warnings": a.warnings,
                }
                for a in analyses
            ],
            "matches": [m.__dict__ for m in matches],
            "rules": rule_results,
            "anomalies": anomalies,
            "clusters": clusters,
            "similarity": similarity,
            "bom": bom_summary,
        }
        path.write_text(json.dumps(payload, indent=2, default=str), encoding="utf-8")
